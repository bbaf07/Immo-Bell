import re
src = open(r"C:\Users\HP\AppData\Local\Temp\ibr\mkpptx.py", encoding="utf8").read().split("# 1. couverture")[0]
exec(src)  # helpers, S, timing, head, foot, card
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_LABEL_POSITION

OUT = r"C:\Users\HP\Documents\Immo Bell\immobell\docs\rapports\Presentation-ImmoBell-Complete.pptx"
CNS = "http://schemas.openxmlformats.org/drawingml/2006/chart"
ANS = "http://schemas.openxmlformats.org/drawingml/2006/main"
DARKGREY = "3A362B"


def lead(s, text, dark=False):
    return s.txt(128, 285, 1664, 100, [dict(t=text, sz=28, c=DMUTED if dark else MUTED, ls=1.05)])


def chart(s, kind, x, y, w, h, cats, series, colors, title, labels=True, legend=False, fmt=None, mx=None):
    cd = CategoryChartData()
    cd.categories = cats
    for name, vals in series:
        cd.add_series(name, vals)
    s.box(x, y, w, h, fill=CARD, line=EDGE, radius=True)  # fond de la carte, fixe
    gf = s.s.shapes.add_chart(kind, P(x + 10), P(y + 10), P(w - 20), P(h - 20), cd)
    ch = gf.chart
    ch.font.size, ch.font.name = Pt(13), FONT
    ch.font.color.rgb = rgb(INK)
    ch.has_title = True
    ch.chart_title.text_frame.text = title
    r = ch.chart_title.text_frame.paragraphs[0].runs[0]
    r.font.size, r.font.bold, r.font.name = Pt(15), True, FONT
    r.font.color.rgb = rgb(INK)
    ch.has_legend = legend
    if legend:
        ch.legend.position, ch.legend.include_in_layout = XL_LEGEND_POSITION.BOTTOM, False
        ch.legend.font.size = Pt(13)
    pl = ch.plots[0]
    if kind == XL_CHART_TYPE.DOUGHNUT:
        for i, c in enumerate(colors):
            pt = pl.series[0].points[i]
            pt.format.fill.solid()
            pt.format.fill.fore_color.rgb = rgb(c)
        pl.has_data_labels = True
        pl.data_labels.font.size, pl.data_labels.font.bold = Pt(16), True
        pl.data_labels.font.color.rgb = rgb(INK)
        pl.data_labels.show_value = True
    else:
        for se, c in zip(pl.series, colors):
            if kind in (XL_CHART_TYPE.LINE_MARKERS,):
                se.format.line.color.rgb = rgb(c)
                se.format.line.width = Pt(3)
                se.smooth = False
                se.marker.format.fill.solid()
                se.marker.format.fill.fore_color.rgb = rgb(c)
                se.marker.format.line.color.rgb = rgb(c)
            else:
                se.format.fill.solid()
                se.format.fill.fore_color.rgb = rgb(c)
        if labels:
            pl.has_data_labels = True
            pl.data_labels.font.size = Pt(13)
            if fmt:
                pl.data_labels.number_format, pl.data_labels.number_format_is_linked = fmt, False
        if hasattr(pl, "gap_width"):
            pl.gap_width = 60
        va = ch.value_axis
        va.has_major_gridlines = True
        va.major_gridlines.format.line.color.rgb = rgb(EDGE)
        va.format.line.fill.background()
        va.tick_labels.font.size = Pt(12)
        if mx:
            va.maximum_scale = mx
        va.minimum_scale = 0
        ch.category_axis.tick_labels.font.size = Pt(12)
        ch.category_axis.format.line.color.rgb = rgb(EDGE)
    # fond transparent (la carte est dessinée dessous)
    cs = ch._chartSpace
    sp = etree.SubElement(cs, f"{{{CNS}}}spPr")
    etree.SubElement(sp, f"{{{ANS}}}noFill")
    ln = etree.SubElement(sp, f"{{{ANS}}}ln")
    etree.SubElement(ln, f"{{{ANS}}}noFill")
    cs.remove(sp)
    chart_el = cs.find(f"{{{CNS}}}chart")
    chart_el.addnext(sp)
    return gf


def row_cards(s, items, y, h, numbered=True, gold_last=False, tsz=32, bsz=24, labels=None):
    n = len(items)
    w = (1664 - (n - 1) * 24) / n
    out = []
    for i, (a, b) in enumerate(items):
        num = (labels[i] if labels else str(i + 1)) if numbered else None
        out.append(card(s, 128 + i * (w + 24), y, w, h, a, b, tsz, bsz, gold=gold_last and i == n - 1, num=num))
    return out


def grid4(s, items, y=400, h=250, gold_last=True):
    out = []
    for i, (a, b) in enumerate(items):
        out.append(card(s, 128 + (i % 2) * 848, y + (i // 2) * (h + 24), 816, h, a, b, 40, 28, gold=gold_last and i == len(items) - 1))
    return out


def seq(s, cards, last_pop=True):
    for k, c in enumerate(cards):
        s.anim("pop" if (last_pop and k == len(cards) - 1) else "rise", c)


def new(bg, notes, trans=PUSH):
    return S(bg, notes, trans)


# ---------------------------------------------------------------- 1 couverture
s = new(DARK, "Introduction : voici l’essentiel pour la direction. Un objectif de 100 contrats et 50 M FCFA, atteint sans publicité payante, grâce à cinq pôles qui s’enchaînent. Les chiffres viennent de la présentation de synthèse ; le site et la démonstration illustrent l’outil.")
logo = s.pic(LOGO_L, 128, 84, 190, 154)
rule = s.box(128, 276, 200, 8, fill=GOLD)
h = s.txt(128, 320, 1560, 400, [dict(t="Un nom de prestige, une gestion d’excellence", sz=120, b=True, c=PALE, ls=0.9)])
sub = s.txt(128, 760, 1500, 120, [dict(t="Gestion locative et services immobiliers à Bonanjo, Douala", sz=44, c=DMUTED)])
dt = s.txt(128, 940, 1200, 40, [dict(t="Présentation aux investisseurs · Octobre 2026", sz=24, c=DMUTED)])
s.anim("fade", logo); s.anim("wipel", rule); s.anim("rise", h); s.anim("rise", sub); s.anim("fade", dt)
s.finish()

# ---------------------------------------------------------------- 2 cinq pôles
s = new(LIGHT, "Immo Bell enchaîne cinq métiers : offrir, faire connaître, sécuriser, suivre et livrer. Les slides suivantes détaillent chacun d’eux. Commençons par l’offre, au cœur du pôle Marketing.")
head(s, "VUE D’ENSEMBLE", "Cinq pôles, une chaîne")
t = lead(s, "Immo Bell enchaîne cinq métiers : offrir, faire connaître, sécuriser, suivre et livrer.")
cs = row_cards(s, [("Marketing", "Définit l’offre, les prix et les leviers"), ("Communication", "Fait connaître sans publicité"), ("Juridique", "Sécurise chaque dossier"), ("Admin & RH", "Suit les baux et le personnel"), ("Transport", "Livre déménagement et mobilité")], 400, 380, gold_last=True)
foot(s); s.anim("fade", t); seq(s, cs); s.finish()

# ---------------------------------------------------------------- 3 offre
s = new(LIGHT, "Le pôle Marketing structure l’entreprise autour d’une activité première, l’immobilier, et de trois activités qui la renforcent. Voyons comment l’immobilier se vend : les trois packs.")
head(s, "1 · L’OFFRE · MARKETING", "Quatre activités")
t = lead(s, "Une activité première, l’immobilier, et trois activités qui la renforcent.")
cs = grid4(s, [("Immobilier", "Activité première : location, meublés, terrains, fonds de commerce."), ("Logistique", "Déménagements, VTC et achat international, par exemple en Chine."), ("Formation hôtelière", "Accueil, restauration, hébergement et gestion."), ("Maintenance", "Suivi préventif et urgences 24 h des appareils électriques.")], y=380, h=250, gold_last=False)
foot(s); s.anim("fade", t); seq(s, cs, last_pop=False); s.finish()

# ---------------------------------------------------------------- 4 packs
s = new(LIGHT, "Trois packs sans remise, pensés par profil de propriétaire. Leur commission se situe dans la fourchette du marché ou en dessous. Reste à transformer ces packs en 100 contrats signés. Le site de démonstration utilise les noms Basic, Standard, Premium et Diaspora : à aligner.")
head(s, "1 · L’OFFRE · MARKETING", "Des packs sous le marché")
t = lead(s, "Trois packs sans remise, pensés par profil de propriétaire. Leur commission se situe dans la fourchette du marché ou en dessous.")
c = chart(s, XL_CHART_TYPE.COLUMN_CLUSTERED, 128, 400, 1000, 540, ["Essentiel", "Confort", "Premium Diaspora"], [("Immo Bell", [8, 12, 18]), ("Marché (bas)", [10, 10, 15]), ("Marché (haut)", [15, 15, 25])], [GOLD, "8A8372", INK], "Commission en % du loyer", legend=True, mx=30)
a = card(s, 1170, 400, 622, 250, "8 %", "Essentiel : quittances, relance impayés J+3", 72, 26)
b = card(s, 1170, 680, 622, 260, "18 % + 25 000 F", "Premium Diaspora : rapport vidéo, visio, bail avocat", 44, 26, gold=True)
foot(s); s.anim("fade", t); s.anim("rise", c); s.anim("rise", a); s.anim("pop", b); s.finish()

# ---------------------------------------------------------------- 5 100 contrats
s = new(LIGHT, "La trajectoire est organique : recensement du Canton Bell, notaires, contenus preuve et parrainage, sans publicité. Pour faire connaître l’offre, place à la communication.")
head(s, "1 · L’OFFRE · MARKETING", "100 contrats en 12 semaines")
t = lead(s, "La trajectoire est organique : recensement du Canton Bell, notaires, contenus preuve et parrainage, sans publicité.")
c = chart(s, XL_CHART_TYPE.LINE_MARKERS, 128, 400, 1000, 540, ["S2 · Recenser", "S6 · Acquisition", "S9 · Montée en gamme", "S12 · Conversion"], [("Contrats cumulés", [0, 30, 50, 100])], [GOLD], "Contrats signés cumulés (objectif)", mx=120)
a = card(s, 1170, 400, 622, 250, "50 M FCFA", "100 contrats × 500 000 F, sans budget publicitaire", 56, 26, gold=True)
b = card(s, 1170, 680, 622, 260, "200 biens", "à recenser en 30 jours, dont 50 % de membres Canton Bell", 56, 26)
foot(s); s.anim("fade", t); s.anim("wipel", c); s.anim("pop", a); s.anim("rise", b); s.finish()

# ---------------------------------------------------------------- 6 réseaux
s = new(LIGHT, "On crée une fois, on republie partout. Trois réseaux principaux, deux en reprise, et des outils gratuits : Canva, CapCut, WhatsApp Business. Mesurons ce que ces réseaux doivent rapporter.")
head(s, "2 · LA VISIBILITÉ · COMMUNICATION", "Cinq réseaux, un contenu")
t = lead(s, "On crée une fois, on republie partout. Trois réseaux principaux, deux en reprise, et des outils gratuits : Canva, CapCut, WhatsApp Business.")
tb = s.table(128, 410, 1664, [["Réseau", "Rôle", "Rythme", "Objectif mois 5"], ["Facebook", "Pages et groupes locaux", "4 posts / semaine", "800 abonnés"], ["TikTok", "Confiance par la vidéo", "3 vidéos / semaine", "600 abonnés"],
                              ["WhatsApp Business", "Répondre, fixer les rendez-vous", "Statut 4 fois / semaine", "140 messages"], ["Instagram", "Vitrine : mêmes contenus", "3 posts / semaine", "300 abonnés"], ["LinkedIn", "Diaspora, cadres, partenaires", "1 post / semaine", "150 abonnés"]], [0.22, 0.34, 0.24, 0.20], 28, "llll")
foot(s); s.anim("fade", t); s.anim("rise", tb); s.finish()

# ---------------------------------------------------------------- 7 entonnoir
s = new(LIGHT, "Sur cinq mois, 25 000 personnes touchées doivent produire 140 messages, dont des contacts sérieux, puis 11 contrats signés. Pour tenir ces chiffres, il faut un pilotage régulier.")
head(s, "2 · LA VISIBILITÉ · COMMUNICATION", "140 messages, 11 contrats")
t = lead(s, "Sur cinq mois, 25 000 personnes touchées doivent produire 140 messages, dont des contacts sérieux, puis 11 contrats signés.")
c = chart(s, XL_CHART_TYPE.BAR_CLUSTERED, 128, 400, 1000, 540, ["Contrats signés", "Rendez-vous ou visites", "Contacts sérieux", "Messages reçus"], [("Volume", [11, 28, 56, 140])], [GOLD], "Entonnoir de conversion sur 5 mois", mx=160)
b = s.box(1170, 400, 622, 540, fill=CARD, line=EDGE, radius=True, pad=30, paras=[dict(t="Comment lire l’entonnoir", sz=32, b=True, after=12),
    dict(t="Message reçu : toute prise de contact (WhatsApp, Facebook, TikTok).", sz=22, c=MUTED, after=6), dict(t="Contact sérieux : propriétaire avec un bien à confier. 40 % des messages (56 sur 140).", sz=22, c=MUTED, after=6),
    dict(t="Rendez-vous ou visite : 50 % des contacts sérieux.", sz=22, c=MUTED, after=6), dict(t="Contrat signé : 40 % des rendez-vous, soit 5,5 M FCFA.", sz=22, c=MUTED)])
foot(s); s.anim("fade", t); s.anim("wipel", c); s.anim("rise", b); s.finish()

# ---------------------------------------------------------------- 8 pilotage
s = new(LIGHT, "Chaque jour a son rendez-vous éditorial : conseil, bien du mardi, preuve, question, offre, coulisses. Trois réunions et des réflexes de crise cadrent le tout. Une fois les contacts générés, il faut sécuriser les dossiers.")
head(s, "2 · LA VISIBILITÉ · COMMUNICATION", "Un pilotage en trois temps")
t = lead(s, "Chaque jour a son rendez-vous éditorial : conseil, bien du mardi, preuve, question, offre, coulisses. Trois réunions et des réflexes de crise cadrent le tout.")
cs = row_cards(s, [("Calendrier", "15 minutes pour fixer les publications de la semaine."), ("Chiffres", "20 minutes sur portée, messages et contacts sérieux."), ("Bilan", "1 heure pour régler le mois suivant."), ("Réflexes", "Rumeur d’arnaque : direction alertée sous 1 h ; erreur corrigée sous 2 h.")], 420, 360, labels=["LUN.", "VEN.", "MOIS", "CRISE"], gold_last=True)
foot(s); s.anim("fade", t); seq(s, cs); s.finish()

# ---------------------------------------------------------------- 9 juridique 3 étapes
s = new(LIGHT, "Le pôle juridique étudie chaque dossier de location, ou de fonds de commerce, avant toute signature. Regardons comment la solvabilité est mesurée.")
head(s, "3 · LA SÉCURITÉ · JURIDIQUE", "Un dossier en trois étapes")
t = lead(s, "Le pôle juridique étudie chaque dossier de location, ou de fonds de commerce, avant toute signature.")
cs = row_cards(s, [("Collecter et vérifier", "Bailleur : CNI et titre de propriété. Demandeur : identité, titre de séjour, avis d’imposition, NIU."), ("Mesurer la solvabilité", "Revenu net, capacité locative, dépôt de garantie, contact de l’ancien bailleur."), ("Conclure", "Visite et état des lieux contradictoire, bail, lettre d’engagement, assurance loyers impayés.")], 400, 400, gold_last=True, tsz=36, bsz=26)
foot(s); s.anim("fade", t); seq(s, cs); s.finish()

# ---------------------------------------------------------------- 10 seuil
s = new(LIGHT, "Règle : le loyer ne doit pas dépasser le tiers du revenu mensuel net. Exemple avec un demandeur à 6 000 000 F de revenus par an. Le simulateur du site applique cette règle automatiquement.")
head(s, "3 · LA SÉCURITÉ · JURIDIQUE", "Un loyer sous le seuil")
t = lead(s, "Règle : le loyer ne doit pas dépasser le tiers du revenu mensuel net. Exemple avec un demandeur à 6 000 000 F de revenus par an.")
c = chart(s, XL_CHART_TYPE.COLUMN_CLUSTERED, 128, 400, 1000, 540, ["Capacité maximale (500 000 ÷ 3)", "Loyer 150 000 F : accepté", "Loyer 200 000 F : refusé"], [("Montant (F)", [166667, 150000, 200000])], [GOLD], "Loyer mensuel (F) : capacité et loyers testés", fmt='# ##0', mx=250000)
b = s.box(1170, 400, 622, 540, fill=CARD, line=EDGE, radius=True, pad=30, paras=[dict(t="Lire l’exemple pas à pas", sz=32, b=True, after=12),
    dict(t="Revenus : 6 000 000 F par an, soit 500 000 F par mois.", sz=22, c=MUTED, after=6), dict(t="Capacité locative : 500 000 ÷ 3 = 166 667 F, le loyer maximal.", sz=22, c=MUTED, after=6),
    dict(t="Loyer de 150 000 F : sous la capacité, dossier accepté.", sz=22, c=MUTED, after=6), dict(t="Loyer de 200 000 F : au-dessus, dossier refusé.", sz=22, c=MUTED, after=6),
    dict(t="Assurance loyers impayés : 2,5 à 4,5 %, soit 3 750 à 6 750 F par mois.", sz=22, c=MUTED)])
foot(s); s.anim("fade", t); s.anim("rise", c); s.anim("rise", b); s.finish()

# ---------------------------------------------------------------- 11 cinq étapes
s = new(LIGHT, "De l’entrée en relation aux cas particuliers, chaque étape a un responsable identifié. Le dossier validé passe ensuite à l’administration. Le site de démonstration montre ce parcours à l’écran.")
head(s, "3 · LA SÉCURITÉ · JURIDIQUE", "Cinq étapes, des rôles clairs")
t = lead(s, "De l’entrée en relation aux cas particuliers, chaque étape a un responsable identifié.")
cs = row_cards(s, [("Entrée en relation", "Analyste juridique : étude du dossier"), ("État des lieux", "Bailleur, agent Immo Bell et locataire"), ("Approbation", "Chef de projet ; commission pour les fonds de commerce"), ("Suivi", "Analyste : contrôle des échéanciers"), ("Cas particuliers", "Décès, impayés, sous-location : huissier, avocat, assurance")], 400, 430, gold_last=True)
foot(s); s.anim("fade", t); seq(s, cs); s.finish()

# ---------------------------------------------------------------- 12 site (plateforme)
s = new(DARK, "Démonstration en direct possible : « Mon espace », puis un clic sur l’un des cinq profils. Les profils de démonstration enregistrent leurs actions dans le navigateur ; ce ne sont pas de vrais comptes. Le serveur réel gère de vrais comptes avec des droits contrôlés côté serveur.")
head(s, "L’OUTIL · LA PLATEFORME", "Un site, cinq espaces", dark=True)
t = lead(s, "Chaque pôle a son espace : le client, le propriétaire, l’analyste juridique, le prestataire et l’administrateur.", dark=True)
data = [("Client", "Cherche, réserve et suit son dossier en 6 étapes."), ("Propriétaire", "Confie ses biens, suit calendrier et relevés."), ("Analyste juridique", "Étudie les dossiers sous 72 h et décide."), ("Prestataire agréé", "Reçoit urgences et devis, rend son rapport."), ("Administrateur", "Valide les biens et pilote tout le reste.")]
cs = []
for i, (a, b) in enumerate(data):
    cs.append(s.box(128 + i * 336, 400, 312, 300, fill=GOLD if i == 4 else "1D1B15", line=None if i == 4 else "3A362B", radius=True, pad=30,
                    paras=[dict(t=a, sz=32, b=True, c=INK if i == 4 else PALE, after=12), dict(t=b, sz=24, c=INK if i == 4 else DMUTED, ls=1.1)]))
f = s.txt(128, 760, 1664, 120, [dict(t="Recherche, visite 360°, simulateur de budget, assistante Awa, français et anglais, thème clair et sombre, version mobile.", sz=30, c=DMUTED)])
foot(s, dark=True); s.anim("fade", t); seq(s, cs); s.anim("fade", f); s.finish()

# ---------------------------------------------------------------- 13 parcours dossier (plateforme)
s = new(LIGHT, "Étapes affichées au client dans son tableau de bord. Le pôle juridique décide sous 72 h ; l’administration voit les délais qui approchent. La solvabilité est testée automatiquement avec la règle du tiers.")
head(s, "L’OUTIL · LA PLATEFORME", "De la demande à la remise des clés")
steps = ["Demande reçue", "Dossier complet", "Étude juridique", "Validée", "État des lieux", "Remise des clés"]
cs = [card(s, 128 + i * 281, 330, 258, 250, st, None, 32, 24, gold=i == 5, num=str(i + 1)) for i, st in enumerate(steps)]
ts = [s.txt(128 + i * 568, 660, 520, 160, [dict(t=[(a, True), (b, False)], sz=32, ls=1.05)]) for i, (a, b) in enumerate([("72 h", " pour décider, avec un compte à rebours visible."), ("Motif obligatoire", " en cas de rejet."), ("Client prévenu", " à chaque étape.")])]
foot(s); seq(s, cs)
for x in ts:
    s.anim("fade", x)
s.finish()

# ---------------------------------------------------------------- 14 sécurité (plateforme)
s = new(DARK, "Mots de passe hachés, sessions par cookie signé, limitation des tentatives de connexion. Ces règles sont vérifiées par des tests automatiques. Un audit indépendant n’a pas encore été fait : ne pas l’affirmer.", FADE)
head(s, "L’OUTIL · LA PLATEFORME", "Les droits sont contrôlés par le serveur", dark=True, two=True)
items = ["Un client ne voit que ses propres dossiers", "Un propriétaire ne valide pas son propre bien", "Un prestataire ne voit que ses interventions", "Seuls l’analyste juridique et l’administrateur décident", "80 tests automatiques sur l’API"]
pairs = [(s.box(128, 434 + i * 80, 16, 16, fill=GOLD), s.txt(176, 410 + i * 80, 1600, 60, [dict(t=it, sz=36, c=PALE)])) for i, it in enumerate(items)]
t = s.txt(128, 830, 1664, 100, [dict(t="À faire avant l’ouverture au public : audit de sécurité indépendant et double authentification des administrateurs.", sz=28, c=DMUTED)])
foot(s, dark=True)
for a, b in pairs:
    s.anim("rise", a, b)
s.anim("fade", t); s.finish()

# ---------------------------------------------------------------- 15 circuit sans blocage
s = new(LIGHT, "Le pôle suit chaque bail : contrôle des pièces, relances, transmission, et escalade dès qu’un point bloque. Détaillons les dossiers et le support RH.")
head(s, "4 · LE SUIVI · ADMINISTRATIF & RH", "Un circuit sans blocage")
t = lead(s, "Le pôle suit chaque bail : contrôle des pièces, relances, transmission, et escalade dès qu’un point bloque.")
cs = row_cards(s, [("Réception et contrôle", "Complétude des pièces, relance si besoin"), ("Suivi et relance", "Baux, attestations et échéances tracés"), ("Transmission", "Information des responsables, suivi des validations"), ("En cas de blocage", "Identifier, coordonner, escalader. Décès, impayés, sous-location : l’analyste juridique traite, avec huissier, avocat et assurance si besoin.")], 400, 480, gold_last=True, bsz=24)
foot(s); s.anim("fade", t); seq(s, cs); s.finish()

# ---------------------------------------------------------------- 16 dossiers et RH
s = new(LIGHT, "Le contrat se clôt par trois signatures : état des lieux, bail, lettre d’engagement et déclaration sur l’honneur. Place maintenant à la livraison des services.")
head(s, "4 · LE SUIVI · ADMINISTRATIF & RH", "Dossiers et support RH")
t = lead(s, "Le contrat se clôt par trois signatures : état des lieux, bail, lettre d’engagement et déclaration sur l’honneur.")
cs = row_cards(s, [("Dossier propriétaire", "Pièces d’identité, documents de propriété, informations sur le bien."), ("Dossier locataire", "Justificatifs, contrats, garanties, assurances et états des lieux à jour."), ("Support au personnel", "Contrats, recrutement, intégration, plan de formation et conformité sociale.")], 400, 400, numbered=True, labels=["BAILLEUR", "LOCATAIRE", "RH"], gold_last=True, tsz=36, bsz=26)
foot(s); s.anim("fade", t); seq(s, cs); s.finish()

# ---------------------------------------------------------------- 17 logistique
s = new(LIGHT, "À Douala 1er, les moyens logistiques ne sont pas intégrés au Canton Bell : le pôle Transport comble ce manque. Voyons les tarifs de ces services.")
head(s, "5 · LA LIVRAISON · TRANSPORT & SUPPLY CHAIN", "Une logistique à intégrer")
t = lead(s, "À Douala 1er, les moyens logistiques ne sont pas intégrés au Canton Bell : le pôle Transport comble ce manque.")
cs = row_cards(s, [("Moyens dispersés", "Pas de service de déménagement professionnel ni de flotte de mobilité urbaine."), ("Base et flotte", "Base de données des moyens, services de déménagement, petite flotte de véhicules."), ("Trois publics", "Riverains de Douala 1er, bailleurs et propriétaires, locataires (cadres, expatriés).")], 400, 400, numbered=True, labels=["CONSTAT", "SOLUTION", "CIBLES"], gold_last=True, tsz=36, bsz=26)
foot(s); s.anim("fade", t); seq(s, cs); s.finish()

# ---------------------------------------------------------------- 18 tarifs
s = new(LIGHT, "Les prix des services annexes sont fixés pour rassurer une clientèle de meublés et de cadres, sans viser le low-cost. Quelle est l’avancée de la roadmap ? Le site propose trois formules de déménagement sans tarif : à aligner avec ces prix.")
head(s, "5 · LA LIVRAISON · TRANSPORT & SUPPLY CHAIN", "Les tarifs à retenir")
t = lead(s, "Les prix des services annexes sont fixés pour rassurer une clientèle de meublés et de cadres, sans viser le low-cost.")
tb = s.table(128, 410, 1664, [["Service", "Tarif", "Repère"], ["VTC", "3 000 F + 300 F/km", "Marché Yango : 2 000 à 3 500 F"], ["Déménagement", "50 k / 90 k / 150 k F", "Studio, 2 chambres, 3 chambres"], ["Achat international", "12 % de commission", "Minimum 50 000 F"],
                              ["Maintenance meublé", "20 000 F par mois", "Visite préventive mensuelle"], ["Intervention urgence", "10 000 F + pièces", "Disponible 24 h"], ["Audit électrique", "30 000 F", "Avant mise en location"]], [0.30, 0.30, 0.40], 28, "lll")
foot(s); s.anim("fade", t); s.anim("rise", tb); s.finish()

# ---------------------------------------------------------------- 19 roadmap
s = new(LIGHT, "La roadmap avance : trois actions sur six sont faites, et le lancement de la flotte est fixé au 1er janvier 2027. Regardons les activités complémentaires : formation, maintenance, mécénat.")
head(s, "5 · LA LIVRAISON · TRANSPORT & SUPPLY CHAIN", "Déjà la moitié du chemin")
t = lead(s, "La roadmap avance : trois actions sur six sont faites, et le lancement de la flotte est fixé au 1er janvier 2027.")
c = chart(s, XL_CHART_TYPE.LINE_MARKERS, 128, 400, 1000, 540, ["07/09", "11/09", "13/09", "14/09", "30/10", "01/01/27"], [("Prévu (6 actions)", [1, 2, 3, 4, 5, 6]), ("Réalisé", [1, 2, 2, 3, None, None])], [INK, GOLD], "Actions cumulées : prévu et réalisé", legend=True, labels=False, mx=7)
b = s.box(1170, 400, 622, 540, fill=CARD, line=EDGE, radius=True, pad=30, paras=[dict(t="Lire la courbe", sz=32, b=True, after=12),
    dict(t="Courbe noire : les 6 actions prévues, cumulées dans le temps.", sz=22, c=MUTED, after=6), dict(t="Courbe or : 3 actions réalisées (logo, formulaire véhicules, descente terrain).", sz=22, c=MUTED, after=6),
    dict(t="En cours : site web et e-mail pro.", sz=22, c=MUTED, after=6), dict(t="À venir : déménagements le 30/10, flotte Yango le 01/01/2027.", sz=22, c=MUTED)])
foot(s); s.anim("fade", t); s.anim("wipel", c); s.anim("rise", b); s.finish()

# ---------------------------------------------------------------- 20 compléments
s = new(LIGHT, "Trois leviers à coût quasi nul renforcent l’offre immobilière et créent des revenus récurrents. Faisons le point sur les moyens à mobiliser.")
head(s, "ACTIVITÉS COMPLÉMENTAIRES", "Formation, maintenance, mécénat")
t = lead(s, "Trois leviers à coût quasi nul renforcent l’offre immobilière et créent des revenus récurrents.")
cs = row_cards(s, [("Formation hôtelière", "Modules de 35 000 à 120 000 F ; 350 000 F de chiffre d’affaires visés dès le mois 2."), ("Maintenance", "20 000 F par mois : 300 000 F de revenu récurrent visé au mois 3."), ("Mécénat", "Nuitées et salles creuses données : 300 000 F de valeur pour moins de 50 000 F de coût.")], 400, 430, numbered=True, labels=["10 / MOIS", "15 ABONNÉS", "3 CONVENTIONS"], gold_last=True, tsz=36, bsz=26)
foot(s); s.anim("fade", t); seq(s, cs); s.finish()

# ---------------------------------------------------------------- 21 moyens
s = new(LIGHT, "Chaque pôle demande peu de ressources, avec 0 F de publicité sur toute la période. Quelques écarts entre pôles sont encore à aligner.")
head(s, "LES MOYENS", "Des moyens modestes")
t = lead(s, "Chaque pôle demande peu de ressources, avec 0 F de publicité sur toute la période.")
tb = s.table(128, 400, 1664, [["Pôle", "Ressources demandées"], ["Marketing", "Aucun budget ; Google Forms et Canva gratuits"], ["Communication", "3 profils créatifs, environ 12 h par semaine"], ["Juridique", "1 analyste juridique en CDD, téléphone et crédit"],
                              ["Administratif & RH", "2 analystes juridiques en CDD, outils d’archivage"], ["Transport", "2 Toyota Yaris, site web, shooting photo, impression"]], [0.28, 0.72], 30, "ll")
foot(s); s.anim("fade", t); s.anim("rise", tb); s.finish()

# ---------------------------------------------------------------- 22 décision
s = new(LIGHT, "Les réseaux ne portent qu’un dixième de l’objectif : le terrain et les partenaires font le reste. Quatre écarts sont à aligner. Moyens : 0 F de publicité, 12 h par semaine, 2 Yaris, analystes en CDD.")
head(s, "LA DÉCISION · DIRECTION", "Ce qu’il faut trancher")
t = lead(s, "Les réseaux ne portent qu’un dixième de l’objectif : le terrain et les partenaires font le reste.")
c = chart(s, XL_CHART_TYPE.DOUGHNUT, 128, 400, 760, 540, ["Réseaux sociaux", "Terrain et partenaires"], [("Contrats", [11, 89])], [INK, GOLD], "Origine visée des 100 contrats", legend=True)
b = s.box(920, 400, 872, 540, fill=GOLD, radius=True, pad=36, paras=[dict(t="À aligner entre pôles", sz=36, b=True, after=14), dict(t="Analystes juridiques : 1 ou 2 ?", sz=28, after=10), dict(t="Dates : 15 septembre ou 15 octobre 2026", sz=28, after=10),
    dict(t="Véhicules : 10 avant l’ouverture ou 15 en 21 jours ?", sz=28, after=10), dict(t="Rythme : 100 contrats en 12 semaines ou 5 mois", sz=28)])
foot(s); s.anim("fade", t); s.anim("wipel", c); s.anim("pop", b); s.finish()

# ---------------------------------------------------------------- 23 écarts de synchronisation
s = new(LIGHT, "Ces écarts apparaissent en comparant la présentation de synthèse avec le site, les rapports et la démonstration. Ils sont à arbitrer entre associés avant la présentation aux investisseurs : aucun n’a été tranché à la place de l’équipe.")
head(s, "SYNCHRONISATION", "Écarts à aligner avant la présentation")
tb = s.table(128, 310, 1664, [["Sujet", "Synthèse de direction", "Site, rapports et démo"],
                              ["Publicité", "0 F sur toute la période", "Rapport n° 2 : 300 000 à 800 000 F par mois de marketing (hypothèse)"],
                              ["Équipe", "1 à 2 analystes en CDD, 3 profils créatifs à 12 h par semaine", "Rapport n° 2 : 11 personnes à l’agence de Bonanjo (hypothèse)"],
                              ["Noms des packs", "Essentiel 8 %, Confort 12 %, Premium Diaspora 18 % + 25 000 F", "Site : Basic 8 %, Standard 12 %, Premium 18 %, Diaspora 18 % + 25 000 F"],
                              ["Déménagement", "50 k, 90 k ou 150 k F (studio, 2 ou 3 chambres)", "Site : formules Classique, Standard, Golf, sans tarif"],
                              ["Activités", "Immobilier, logistique, formation hôtelière, maintenance", "Site : pas de formation ni de mécénat ; entretien par prestataires agréés"]], [0.16, 0.42, 0.42], 22, "lll")
foot(s); s.anim("rise", tb); s.finish()

# ---------------------------------------------------------------- 24 deux échelles
s = new(LIGHT, "Deux niveaux de lancement coexistent dans les documents. La synthèse de direction décrit un lancement léger (0 F de publicité, quelques personnes). Le rapport n° 2 chiffre une agence complète avec des hypothèses que j’ai estimées moi-même, sans devis. À arbitrer entre associés : quel périmètre présenter aux investisseurs.")
head(s, "SYNCHRONISATION", "Deux échelles de lancement")
a = s.box(128, 330, 816, 530, fill=CARD, line=EDGE, radius=True, pad=40, paras=[dict(t="LANCEMENT LÉGER", sz=24, b=True, c=EYEB, after=10), dict(t="Synthèse de direction", sz=40, b=True, after=16),
    dict(t="0 F de publicité, outils gratuits", sz=28, c=MUTED, after=8), dict(t="1 à 2 analystes en CDD, 3 profils créatifs", sz=28, c=MUTED, after=8), dict(t="2 Toyota Yaris", sz=28, c=MUTED, after=8), dict(t="100 contrats et 50 M FCFA en 12 semaines", sz=28, c=MUTED)])
b = s.box(976, 330, 816, 530, fill=GOLD, radius=True, pad=40, paras=[dict(t="AGENCE COMPLÈTE", sz=24, b=True, after=10), dict(t="Rapport n° 2 (hypothèses)", sz=40, b=True, after=16),
    dict(t="Agence à Bonanjo, 11 personnes, plus 7 en logistique", sz=28, after=8), dict(t="Besoin de 49 à 112 M FCFA, puis 11 à 26 M pour la logistique", sz=28, after=8), dict(t="Seuil de rentabilité : 98 à 200 biens loués", sz=28, after=8), dict(t="Chiffres estimés, à remplacer par des devis locaux", sz=28)])
n = s.txt(128, 880, 1664, 60, [dict(t="À arbitrer entre associés : quel périmètre présenter aux investisseurs.", sz=28, c=MUTED)])
foot(s); s.anim("rise", a); s.anim("pop", b); s.anim("fade", n); s.finish()

# ---------------------------------------------------------------- 25 merci
s = new(DARK, "Remercier les investisseurs pour leur temps, puis ouvrir les questions et lancer la démonstration. Garder sous la main le rapport n° 2 pour les chiffres et le guide de démonstration pour les comptes fictifs.", None)
logo = s.pic(LOGO_L, 128, 84, 190, 154)
m = s.txt(128, 340, 1664, 220, [dict(t="Merci", sz=160, b=True, c=PALE)])
rule = s.box(128, 570, 200, 8, fill=GOLD)
q = s.txt(128, 620, 1500, 120, [dict(t="Place à vos questions et à la démonstration en direct.", sz=44, c=DMUTED)])
c = s.txt(128, 850, 1500, 130, [dict(t="contact@immobell.cm", sz=32, c=PALE), dict(t="Bonanjo, Douala · Service après-vente de 8 h à 18 h", sz=32, c=PALE)])
s.anim("fade", logo); s.anim("rise", m); s.anim("wipel", rule); s.anim("rise", q); s.anim("fade", c)
s.finish()

prs.save(OUT)
print("ok", len(prs.slides._sldIdLst))
