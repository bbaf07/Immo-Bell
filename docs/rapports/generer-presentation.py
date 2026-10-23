from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from lxml import etree

PUB = r"C:\Users\HP\Documents\Immo Bell\immobell\web\public"
OUT = r"C:\Users\HP\Documents\Immo Bell\immobell\docs\rapports\Presentation-ImmoBell.pptx"
FONT = "Arial"
LIGHT, DARK = "F4EFE2", "0F0E0B"
INK, MUTED, EYEB, CARD, EDGE, GOLD, NUM, DMUTED, PALE = "1A1813", "5A5547", "7A5A00", "FBF8EF", "DCD4BD", "F2B000", "C28F00", "BDB6A3", "F4EFE2"
P = lambda v: Emu(int(v * 6350))
rgb = lambda h: RGBColor.from_string(h)

prs = Presentation()
prs.slide_width, prs.slide_height = Emu(12192000), Emu(6858000)
BLANK = prs.slide_layouts[6]


class S:
    def __init__(self, bg, notes, trans):
        self.s = prs.slides.add_slide(BLANK)
        self.s.background.fill.solid()
        self.s.background.fill.fore_color.rgb = rgb(bg)
        self.s.notes_slide.notes_text_frame.text = notes
        self.trans, self.groups, self.bg = trans, [], bg

    def box(self, x, y, w, h, fill=None, line=None, radius=False, paras=None, pad=0, anchor="t"):
        shp = self.s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE, P(x), P(y), P(w), P(h))
        if radius:
            shp.adjustments[0] = min(0.5, 16 / min(w, h))
        if fill:
            shp.fill.solid()
            shp.fill.fore_color.rgb = rgb(fill)
        else:
            shp.fill.background()
        if line:
            shp.line.color.rgb = rgb(line)
            shp.line.width = Pt(0.75)
        else:
            shp.line.fill.background()
        shp.shadow.inherit = False
        tf = shp.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = P(pad)
        tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE}[anchor]
        for i, d in enumerate(paras or []):
            para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            para.alignment = PP_ALIGN.LEFT
            para.space_after = Pt(d.get("after", 0) * 0.5)
            if "ls" in d:
                para.line_spacing = d["ls"]
            runs = d["t"] if isinstance(d["t"], list) else [(d["t"], d.get("b", False))]
            for txt, bold in runs:
                r = para.add_run()
                r.text = txt
                r.font.name = FONT
                r.font.size = Pt(d["sz"] * 0.5)
                r.font.bold = bold
                r.font.color.rgb = rgb(d.get("c", INK))
        return shp

    def txt(self, x, y, w, h, paras):
        return self.box(x, y, w, h, paras=paras)

    def pic(self, path, x, y, w, h):
        return self.s.shapes.add_picture(path, P(x), P(y), P(w), P(h))

    def table(self, x, y, w, rows, widths, sz, aligns, bold_rows=()):
        gf = self.s.shapes.add_table(len(rows), len(rows[0]), P(x), P(y), P(w), P(sz * 2.1 * len(rows)))
        tb = gf.table
        for ci, cw in enumerate(widths):
            tb.columns[ci].width = P(w * cw)
        for ri, row in enumerate(rows):
            tb.rows[ri].height = P(sz * 2.1)
            for ci, val in enumerate(row):
                c = tb.cell(ri, ci)
                c.fill.solid()
                c.fill.fore_color.rgb = rgb(INK if ri == 0 else (CARD if ri % 2 else "F0EADA"))
                c.vertical_anchor = MSO_ANCHOR.MIDDLE
                p = c.text_frame.paragraphs[0]
                p.alignment = PP_ALIGN.RIGHT if aligns[ci] == "r" else PP_ALIGN.LEFT
                r = p.add_run()
                r.text = val
                r.font.name, r.font.size = FONT, Pt(sz * 0.5)
                r.font.bold = ri == 0 or ri in bold_rows
                r.font.color.rgb = rgb(PALE if ri == 0 else INK)
        return gf

    def anim(self, effect, *shapes):
        self.groups.append([(effect, sh) for sh in shapes])

    def finish(self):
        el = self.s._element
        if self.trans:
            tr = etree.fromstring(f'<p:transition xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" spd="med">{self.trans}</p:transition>')
            el.append(tr)
        if self.groups:
            el.append(timing(self.groups))


NS = 'xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"'
DUR = 500


def effect_xml(kind, spid, ids, node, grp):
    n = lambda: next(ids)
    pid, sub = {"fade": (10, 0), "rise": (22, 4), "wipel": (22, 8), "pop": (53, 16)}[kind]
    g = ' grpId="0"' if grp else ""
    tgt = f'<p:tgtEl><p:spTgt spid="{spid}"/></p:tgtEl>'
    vis = f'<p:set><p:cBhvr><p:cTn id="{n()}" dur="1" fill="hold"><p:stCondLst><p:cond delay="0"/></p:stCondLst></p:cTn>{tgt}<p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst></p:cBhvr><p:to><p:strVal val="visible"/></p:to></p:set>'
    if kind == "pop":
        body = "".join(
            f'<p:anim calcmode="lin" valueType="num"><p:cBhvr><p:cTn id="{n()}" dur="{DUR}" fill="hold"/>{tgt}<p:attrNameLst><p:attrName>{a}</p:attrName></p:attrNameLst></p:cBhvr><p:tavLst><p:tav tm="0"><p:val><p:fltVal val="0"/></p:val></p:tav><p:tav tm="100000"><p:val><p:strVal val="#{a}"/></p:val></p:tav></p:tavLst></p:anim>'
            for a in ("ppt_w", "ppt_h"))
        body += f'<p:animEffect transition="in" filter="fade"><p:cBhvr><p:cTn id="{n()}" dur="{DUR}"/>{tgt}</p:cBhvr></p:animEffect>'
    else:
        filt = {"fade": "fade", "rise": "wipe(down)", "wipel": "wipe(left)"}[kind]
        body = f'<p:animEffect transition="in" filter="{filt}"><p:cBhvr><p:cTn id="{n()}" dur="{DUR}"/>{tgt}</p:cBhvr></p:animEffect>'
    return f'<p:par><p:cTn id="{n()}" presetID="{pid}" presetClass="entr" presetSubtype="{sub}" fill="hold"{g} nodeType="{node}"><p:stCondLst><p:cond delay="0"/></p:stCondLst><p:childTnLst>{vis}{body}</p:childTnLst></p:cTn></p:par>'


def timing(groups):
    counter = iter(range(3, 10000))
    ids = counter
    outer = next(ids)
    inner, cum, bld = "", 0, []
    for grp in groups:
        gid = next(ids)
        eff = ""
        for k, (kind, sh) in enumerate(grp):
            is_sp = sh._element.tag.endswith("}sp")
            if is_sp:
                bld.append(f'<p:bldP spid="{sh.shape_id}" grpId="0" animBg="1"/>')
            eff += effect_xml(kind, sh.shape_id, ids, "afterEffect" if k == 0 else "withEffect", is_sp)
        inner += f'<p:par><p:cTn id="{gid}" fill="hold"><p:stCondLst><p:cond delay="{cum}"/></p:stCondLst><p:childTnLst>{eff}</p:childTnLst></p:cTn></p:par>'
        cum += DUR
    xml = (f'<p:timing {NS}><p:tnLst><p:par><p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot"><p:childTnLst>'
           f'<p:seq concurrent="1" nextAc="seek"><p:cTn id="2" dur="indefinite" nodeType="mainSeq"><p:childTnLst>'
           f'<p:par><p:cTn id="{outer}" fill="hold"><p:stCondLst><p:cond delay="indefinite"/><p:cond evt="onBegin" delay="0"><p:tn val="2"/></p:cond></p:stCondLst><p:childTnLst>{inner}</p:childTnLst></p:cTn></p:par>'
           f'</p:childTnLst></p:cTn><p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst>'
           f'<p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst></p:seq>'
           f'</p:childTnLst></p:cTn></p:par></p:tnLst><p:bldLst>{"".join(bld)}</p:bldLst></p:timing>')
    return etree.fromstring(xml)


PUSH, FADE = '<p:push dir="u"/>', "<p:fade/>"
LOGO_D, LOGO_L = PUB + r"\logo.png", PUB + r"\logo-light.png"


def head(s, eyebrow, title, dark=False, two=False):
    s.txt(128, 120, 1664, 40, [dict(t=eyebrow, sz=24, b=True, c=GOLD if dark else EYEB)])
    s.txt(128, 168, 1664, 170 if two else 90, [dict(t=title, sz=72, b=True, c=PALE if dark else INK, ls=0.95)])


def foot(s, dark=False):
    s.txt(128, 975, 800, 40, [dict(t="Immo Bell · Confidentiel", sz=24, c=DMUTED if dark else MUTED)])
    s.pic(LOGO_L if dark else LOGO_D, 1672, 948, 120, 97)


def card(s, x, y, w, h, title, body, tsz=32, bsz=24, gold=False, num=None):
    paras = []
    if num:
        paras.append(dict(t=num, sz=72, b=True, c=INK if gold else NUM, after=10))
    paras.append(dict(t=title, sz=tsz, b=True, c=INK, after=12))
    if body:
        paras.append(dict(t=body, sz=bsz, c=INK if gold else MUTED, ls=1.1))
    return s.box(x, y, w, h, fill=GOLD if gold else CARD, line=None if gold else EDGE, radius=True, paras=paras, pad=34)


# 1. couverture
s = S(DARK, "Bienvenue. En dix minutes : ce que fait Immo Bell, comment le site protège chaque dossier, puis ce qu’il faut mettre en place à Douala (équipe, budget, calendrier). Les montants sont des hypothèses de travail à valider par devis locaux.", PUSH)
logo = s.pic(LOGO_L, 128, 84, 190, 154)
rule = s.box(128, 276, 200, 8, fill=GOLD)
h = s.txt(128, 320, 1560, 400, [dict(t="Un nom de prestige, une gestion d’excellence", sz=120, b=True, c=PALE, ls=0.9)])
sub = s.txt(128, 760, 1500, 120, [dict(t="Gestion locative et services immobiliers à Bonanjo, Douala", sz=44, c=DMUTED)])
dt = s.txt(128, 940, 1200, 40, [dict(t="Présentation aux investisseurs · Octobre 2026", sz=24, c=DMUTED)])
s.anim("fade", logo); s.anim("wipel", rule); s.anim("rise", h); s.anim("rise", sub); s.anim("fade", dt)
s.finish()

# 2. offre
s = S(LIGHT, "Texte d’accueil du site : nous accompagnons les personnes physiques et morales dans la gestion des biens meubles et immeubles, à travers la gestion locative 360 degrés, les fonds de commerce, les aménagements et déménagements, et l’entretien de l’appareillage et des biens. Immo Bell n’est pas propriétaire des biens : elle étudie les dossiers et gère les demandes.", PUSH)
head(s, "L’OFFRE", "Gérer les biens meubles et immeubles")
t = s.txt(128, 290, 1664, 50, [dict(t="Pour les personnes physiques et morales, à Douala.", sz=32, c=MUTED)])
cs = [card(s, 128, 370, 816, 270, "Gestion locative 360°", "Annonces validées, étude des dossiers sous 72 h, état des lieux, suivi des urgences.", 40, 28),
      card(s, 976, 370, 816, 270, "Fonds de commerce", "Étude des deux parties, évaluation, contrat de vente ou de location-gérance.", 40, 28),
      card(s, 128, 660, 816, 270, "Aménagements et déménagements", "Trois formules, Classique, Standard et Golf, du transport à l’assainissement.", 40, 28),
      card(s, 976, 660, 816, 270, "Entretien des biens", "Plomberie, électricité, jardinage et BTP par des prestataires agréés.", 40, 28, gold=True)]
foot(s)
s.anim("fade", t); s.anim("rise", cs[0]); s.anim("rise", cs[1]); s.anim("rise", cs[2]); s.anim("pop", cs[3])
s.finish()

# 3. site
s = S(LIGHT, "Démonstration en direct possible : « Mon espace », puis un clic sur l’un des cinq profils. Les profils de démonstration enregistrent leurs actions dans le navigateur ; ce ne sont pas de vrais comptes. Le serveur réel gère de vrais comptes avec des droits contrôlés côté serveur.", PUSH)
head(s, "LE SITE", "Un site, cinq espaces")
data = [("Client", "Cherche, réserve et suit son dossier en 6 étapes."), ("Propriétaire", "Confie ses biens, suit calendrier et relevés."), ("Analyste juridique", "Étudie les dossiers sous 72 h et décide."),
        ("Prestataire agréé", "Reçoit urgences et devis, rend son rapport."), ("Administrateur", "Valide les biens et pilote tout le reste.")]
cs = [card(s, 128 + i * 336, 360, 312, 290, a, b, 32, 24, gold=i == 4) for i, (a, b) in enumerate(data)]
t = s.txt(128, 720, 1664, 120, [dict(t="Recherche, visite 360°, simulateur de budget, assistante Awa, français et anglais, thème clair et sombre, version mobile.", sz=32, c=MUTED)])
foot(s)
for c in cs[:4]:
    s.anim("rise", c)
s.anim("pop", cs[4]); s.anim("fade", t)
s.finish()

# 4. parcours
s = S(LIGHT, "Étapes affichées au client dans son tableau de bord. Le pôle juridique décide sous 72 h ; l’administration voit les délais qui approchent. La solvabilité est testée automatiquement avec la règle du tiers : le loyer ne dépasse pas un tiers du revenu net.", PUSH)
head(s, "LE PARCOURS D’UN DOSSIER", "De la demande à la remise des clés")
steps = ["Demande reçue", "Dossier complet", "Étude juridique", "Validée", "État des lieux", "Remise des clés"]
cs = [card(s, 128 + i * 281, 330, 258, 250, st, None, 32, 24, gold=i == 5, num=str(i + 1)) for i, st in enumerate(steps)]
ts = [s.txt(128 + i * 568, 660, 520, 160, [dict(t=[(a, True), (b, False)], sz=32, ls=1.05)]) for i, (a, b) in enumerate([("72 h", " pour décider, avec un compte à rebours visible."), ("Motif obligatoire", " en cas de rejet."), ("Client prévenu", " à chaque étape.")])]
foot(s)
for c in cs[:5]:
    s.anim("rise", c)
s.anim("pop", cs[5])
for t in ts:
    s.anim("fade", t)
s.finish()

# 5. sécurité
s = S(DARK, "Mots de passe hachés, sessions par cookie signé, limitation des tentatives de connexion, lien de réinitialisation à usage unique valable une heure. Ces règles sont vérifiées par des tests automatiques. Un audit indépendant n’a pas encore été fait : ne pas l’affirmer.", FADE)
head(s, "LA SÉCURITÉ", "Les droits sont contrôlés par le serveur", dark=True, two=True)
items = ["Un client ne voit que ses propres dossiers", "Un propriétaire ne valide pas son propre bien", "Un prestataire ne voit que ses interventions", "Seuls l’analyste juridique et l’administrateur décident", "80 tests automatiques sur l’API"]
pairs = [(s.box(128, 434 + i * 80, 16, 16, fill=GOLD), s.txt(176, 410 + i * 80, 1600, 60, [dict(t=it, sz=36, c=PALE)])) for i, it in enumerate(items)]
t = s.txt(128, 830, 1664, 100, [dict(t="À faire avant l’ouverture au public : audit de sécurité indépendant et double authentification des administrateurs.", sz=28, c=DMUTED)])
foot(s, dark=True)
for a, b in pairs:
    s.anim("rise", a, b)
s.anim("fade", t)
s.finish()

# 6. modèle
s = S(LIGHT, "Taux repris du prototype ; la base de calcul définitive des packs reste à arrêter entre associés. Le seuil de rentabilité vient d’un loyer moyen de 425 000 FCFA (moyenne des huit locations d’exemple) et des charges mensuelles estimées dans le rapport n° 2.", PUSH)
head(s, "LE MODÈLE ÉCONOMIQUE", "Une commission sur chaque loyer géré")
tb = s.table(128, 340, 860, [["Formule", "Commission"], ["Basic", "8 % du loyer"], ["Standard", "12 % du loyer"], ["Premium", "18 % du loyer"], ["Diaspora", "18 % + 25 000 FCFA"]], [0.5, 0.5], 32, "ll")
g = s.box(1036, 340, 756, 340, fill=GOLD, radius=True, anchor="m", pad=44, paras=[dict(t="98 à 200 biens loués", sz=44, b=True, after=16), dict(t="pour couvrir les charges mensuelles de la phase 1 (hypothèse, commission moyenne de 12 %).", sz=28, ls=1.1)])
t = s.txt(128, 740, 1664, 80, [dict(t="Services, logistique et flotte s’ajoutent : tarifs à définir avec l’équipe.", sz=28, c=MUTED)])
foot(s)
s.anim("rise", tb); s.anim("pop", g); s.anim("fade", t)
s.finish()

# 7. équipe
s = S(LIGHT, "Salaires bruts mensuels estimés par poste, avec 15 à 20 % de charges patronales. À remplacer par les salaires réels du marché de Douala. L’analyste juridique est appuyé par un avocat conseil externe.", PUSH)
head(s, "L’ÉQUIPE", "Onze personnes à l’ouverture, sept de plus pour la logistique", two=True)
a = card(s, 128, 400, 816, 410, "Agence de Bonanjo", "Direction, exploitation, pôle juridique, conseillers, agents terrain, comptabilité, service après-vente, plateforme, marketing.", 32, 28, num="11")
b = card(s, 976, 400, 816, 410, "Logistique", "Un responsable, quatre déménageurs et deux chauffeurs, à l’ouverture des déménagements.", 32, 28, gold=True, num="7")
t = s.txt(128, 840, 1480, 90, [dict(t="Masse salariale de la phase 1, charges comprises : 3,2 à 5,9 millions de FCFA par mois (hypothèse).", sz=28, c=MUTED)])
foot(s)
s.anim("rise", a); s.anim("pop", b); s.anim("fade", t)
s.finish()

# 8. budget
s = S(LIGHT, "Chiffres du rapport n° 2, section Budget. Fourchettes que j’ai estimées moi-même, ce ne sont pas des devis. Le montant à lever, la répartition entre apport des associés et financement externe, et la rémunération des fondateurs restent à arrêter entre associés.", PUSH)
head(s, "LE BUDGET", "Besoin de financement, en millions de FCFA", two=True)
tb = s.table(128, 410, 1664, [["Poste", "Hypothèse basse", "Hypothèse haute"], ["Investissements de démarrage", "19,3", "50,7"], ["Trésorerie : six mois de fonctionnement", "30,1", "61,1"],
                              ["Lancement et pilote", "49,4", "111,8"], ["Logistique, phase 2", "11,2", "25,7"], ["Besoin total", "60,6", "137,5"]], [0.52, 0.24, 0.24], 32, "lrr", bold_rows=(3, 5))
t = s.txt(128, 850, 1480, 80, [dict(t="Hypothèses de travail, à remplacer par des devis locaux : loyer de Bonanjo, véhicules, salaires, aménagement.", sz=28, c=MUTED)])
foot(s)
s.anim("rise", tb); s.anim("fade", t)
s.finish()

# 9. calendrier
s = S(LIGHT, "Calendrier indicatif. Le site annonce les déménagements dès le 30 octobre 2026 et la flotte Yango au 1er janvier 2027 : ces dates sont à ajuster avec l’équipe selon la mise en place réelle.", FADE)
head(s, "LE CALENDRIER", "Quatre phases sur douze mois")
ph = [("MOIS 0 À 2", "Préparation", "Société, local, recrutement de la direction, premiers prestataires, photos réelles."), ("MOIS 3 À 5", "Pilote Bonanjo et Akwa", "20 à 50 biens, dossiers sous 72 h, urgences avec 3 à 5 prestataires."),
      ("MOIS 6 À 9", "Extension", "Déménagements, assainissement, nouveaux quartiers, paiement Mobile Money."), ("MOIS 10 À 12", "Diaspora et flotte", "Offre Diaspora, flotte Yango, signature électronique et états des lieux numériques.")]
cs = []
for i, (m, t_, d) in enumerate(ph):
    gold = i == 3
    cs.append(s.box(128 + i * 424, 340, 392, 420, fill=GOLD if gold else CARD, line=None if gold else EDGE, radius=True, pad=32,
                    paras=[dict(t=m, sz=24, b=True, c=INK if gold else EYEB, after=12), dict(t=t_, sz=32, b=True, after=12), dict(t=d, sz=24, c=INK if gold else MUTED, ls=1.1)]))
foot(s)
for c in cs[:3]:
    s.anim("rise", c)
s.anim("pop", cs[3])
s.finish()

# 10. suite
s = S(DARK, "Passer à la démonstration : « Mon espace », puis les profils Client, Analyste juridique, Propriétaire, Prestataire et Administrateur. Le guide de démonstration PDF donne le déroulé et les comptes fictifs. Rappeler que les profils sont des profils de démonstration et que les montants sont des hypothèses.", FADE)
head(s, "LA SUITE", "Ce qui reste avant l’ouverture", dark=True)
todo = ["Hébergement, domaine et sécurité", "E-mails et SMS réels", "Paiement Mobile Money", "Photos réelles des biens", "Relecture juridique des contrats", "Étude de marché locale"]
ts = [s.txt(128 + (i % 2) * 864, 340 + (i // 2) * 100, 800, 70, [dict(t=x, sz=32, c=PALE)]) for i, x in enumerate(todo)]
cta = s.txt(128, 800, 1664, 80, [dict(t="Place à la démonstration en direct", sz=44, b=True, c=GOLD)])
foot(s, dark=True)
for t in ts:
    s.anim("fade", t)
s.anim("pop", cta)
s.finish()

# 11. merci
s = S(DARK, "Remercier les investisseurs pour leur temps, puis ouvrir les questions et lancer la démonstration. Garder sous la main le rapport n° 2 pour les chiffres et le guide de démonstration pour les comptes fictifs.", None)
logo = s.pic(LOGO_L, 128, 84, 190, 154)
m = s.txt(128, 340, 1664, 220, [dict(t="Merci", sz=160, b=True, c=PALE)])
rule = s.box(128, 570, 200, 8, fill=GOLD)
q = s.txt(128, 620, 1500, 120, [dict(t="Place à vos questions et à la démonstration en direct.", sz=44, c=DMUTED)])
c = s.txt(128, 850, 1500, 130, [dict(t="contact@immobell.cm", sz=32, c=PALE), dict(t="Bonanjo, Douala · Service après-vente de 8 h à 18 h", sz=32, c=PALE)])
s.anim("fade", logo); s.anim("rise", m); s.anim("wipel", rule); s.anim("rise", q); s.anim("fade", c)
s.finish()

prs.save(OUT)
print("ok", OUT, len(prs.slides._sldIdLst))
